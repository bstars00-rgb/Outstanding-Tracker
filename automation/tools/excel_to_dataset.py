"""
Convert the weekly OP outstanding workbook into a tracker ReceivablesDataset JSON
so it can be run through the pipeline with DATA_SOURCE=file (manual run mode).

Usage:
  python automation/tools/excel_to_dataset.py <workbook.xlsx> <as_of YYYY-MM-DD> [out.json]

Mapping (see docs/RUNBOOK_MANUAL_RUN.md):
  Outstanding sheet     -> invoices (+ payments/activities parsed from "Remark Detail (Ellis)")
                           The sheet is either "Outstanding" or a dated tab such as "02-Oct" (the tab matching <as_of>,
                           otherwise the latest dated tab). The previous dated tab is used as last week's reference.
  "Noted" column        -> "Payment received ..." = money is at the bank but ELLIS is not updated yet (AC team pending).
                           Those invoices are treated as collected in the tracker and queued in the ELLIS reflection
                           chain at stage RECEIVED (recorded_at = null); the ELLIS-ledger figures are kept in the notes.
  Tier sheet            -> customers (PIC(OMH) = account owner, Tier); automation/config/pic-overrides.json wins when present
  information sheet     -> FX table (Rate to JPY)
  automation/input/ellis-channels.json (ELLIS MCP list_channels, optional) -> managing entity + country by seller code
Nothing is invented: fields no source carries are null / 'Unknown' and listed in completeness.notes.
"""
import json
import os
import re
import statistics
import sys
from collections import defaultdict
from datetime import datetime, date

import openpyxl

# Internal accounts (confirmed by Global Ops 2026-09-28): not customer receivables, so they are EXCLUDED from the
# tracker dataset and listed with their balances in completeness.notes instead.
#   "Business Trip in ..."  = employee business trips (출장)
#   "Unsold room (JP)"      = hard-block unsold inventory (하드블럭 미판매분)
INTERNAL_HINTS = ("business trip", "unsold room")

# Managing entity (control company). Primary source: ELLIS list_channels ownerCompCode (by seller code).
# Fallback when the channel file has no row: mapping confirmed by Global Ops 2026-09-28.
ENTITY_BY_OWNER_CODE = {110000: "OMH Korea", 120000: "OMH Japan", 130000: "OMH Vietnam", 160000: "OMH Singapore"}
CONTROL_COMPANY_DEFAULT = "OMH Singapore"
CONTROL_COMPANY_BY_CHANNEL = {"ctrip vietnam": "OMH Vietnam", "ctrip korea": "OMH Korea", "agoda": "OMH Korea", "kakao": "OMH Korea"}

CHANNELS_FILE = "automation/input/ellis-channels.json"
PIC_OVERRIDES_FILE = "automation/config/pic-overrides.json"
NOTE_HDR = re.compile(r"^\[(\d{4}-\d{2}-\d{2})\]\s*(.+?)\s*\|\s*([\d,\.]+)\s*$")
DATED_TAB = re.compile(r"^(\d{1,2})[-\s]([A-Za-z]{3})$")


def iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return None


def num(s):
    return float(str(s).replace(",", ""))


def tab_date(title, as_of):
    """'02-Oct' -> ISO date in the year of as_of (previous year if that would be in the future)."""
    m = DATED_TAB.match(title.strip())
    if not m:
        return None
    try:
        d = datetime.strptime(f"{m.group(1)}-{m.group(2)}-{as_of[:4]}", "%d-%b-%Y").date()
    except ValueError:
        return None
    if d.isoformat() > as_of:
        d = d.replace(year=d.year - 1)
    return d.isoformat()


def pick_sheets(wb, as_of):
    dated = sorted((tab_date(ws.title, as_of), ws.title) for ws in wb.worksheets if tab_date(ws.title, as_of))
    if "Outstanding" in wb.sheetnames:
        return "Outstanding", (dated[-1][1] if dated else None)
    if not dated:
        raise SystemExit("No 'Outstanding' sheet and no dated tab (e.g. '02-Oct') found")
    cur = next((t for d, t in dated if d == as_of), dated[-1][1])
    cur_date = tab_date(cur, as_of)
    prev = [t for d, t in dated if d < cur_date]
    return cur, (prev[-1] if prev else None)


def read_rows(ws):
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    return col, [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]


def load_json(path):
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return None


def main(path, as_of, out):
    wb = openpyxl.load_workbook(path, data_only=True)
    cur_tab, prev_tab = pick_sheets(wb, as_of)
    notes = [f"Manual run: converted from OP workbook {os.path.basename(path)}, sheet '{cur_tab}' (as of {as_of})."]

    # ---- FX (information sheet) ----
    rates = []
    for r in wb["information"].iter_rows(values_only=True):
        if r and isinstance(r[0], str) and r[0].strip().upper() in ("USD", "KRW", "VND", "TWD", "THB", "SGD", "CNY", "HKD", "EUR") and isinstance(r[1], (int, float)):
            rates.append({"currency": r[0].strip().upper(), "rate_to_reporting": float(r[1]), "rate_date": as_of, "source": "OP workbook 'information' sheet (approx. rate to JPY)"})
    fx = {"reporting_currency": "JPY", "as_of": as_of, "rates": rates}
    rate_of = lambda ccy: 1.0 if ccy == "JPY" else next((x["rate_to_reporting"] for x in rates if x["currency"] == ccy), None)
    notes.append("FX: approximate rates from the workbook 'information' sheet, not ELLIS applied rates.")

    # ---- Customers (Tier sheet + overrides + ELLIS channel directory) ----
    tier_by_name = {}
    for r in wb["Tier"].iter_rows(min_row=3, values_only=True):
        if r and r[3]:
            tier_by_name[str(r[3]).strip().lower()] = {"status": r[1], "type": r[2], "channel": str(r[3]).strip(), "company": r[4], "pic": r[5], "tier": r[6]}
    overrides = load_json(PIC_OVERRIDES_FILE) or {}
    pic_override = {k.strip().lower(): v for k, v in (overrides.get("owners") or {}).items()}
    channels = load_json(CHANNELS_FILE)
    chan_by_code = {}
    if channels:
        f = channels["fields"]
        for row in channels["channels"]:
            rec = dict(zip(f, row))
            chan_by_code[str(rec["sellerCompCode"])] = rec

    # ---- Previous week's sheet (reference for payment dating) ----
    prev = {}
    if prev_tab and prev_tab != cur_tab:
        pcol, prows = read_rows(wb[prev_tab])
        for r in prows:
            remark = r[pcol["Remark Detail (Ellis)"]] if "Remark Detail (Ellis)" in pcol else None
            dates = re.findall(r"^\[(\d{4}-\d{2}-\d{2})\]", remark, flags=re.M) if isinstance(remark, str) else []
            prev[str(r[pcol["Invocie No"]]).strip()] = {"balance": float(r[pcol["Balance"]] or 0), "paid": float(r[pcol["Paid amount"]] or 0), "last_payment": max(dates) if dates else None}

    # ---- Invoices ----
    col, inv_rows = read_rows(wb[cur_tab])
    customers = {}
    invoices, payments, activities = [], [], []
    internal_rows = []
    terms = defaultdict(list)
    ccy_count = defaultdict(lambda: defaultdict(int))
    unmatched_tier, pic_changed, entity_from_ellis, entity_fallback = set(), set(), set(), set()
    ledger_jpy = pending_jpy = reopen_jpy = 0.0
    ledger_n = pending_n = reopen_n = 0

    for r in inv_rows:
        inv_no = str(r[col["Invocie No"]]).strip()
        code = str(r[col["Seller code"]]).strip()
        name = str(r[col["Seller name"]]).strip()
        cid = f"seller:{code}"
        ccy = str(r[col["B.Cur"]]).strip().upper()
        total = float(r[col["B.Sum amount"]] or 0)
        paid = float(r[col["Paid amount"]] or 0)
        bal = float(r[col["Balance"]] or 0)
        issue = iso(r[col["Issue date"]])
        due = iso(r[col["Due date"]])
        tier_cell = r[col["Tier"]]
        remark = r[col["Remark Detail (Ellis)"]] if "Remark Detail (Ellis)" in col else None
        vnop = r[col["VNOP action"]] if "VNOP action" in col else None
        noted = r[col["Noted"]] if "Noted" in col else None
        noted = noted.strip() if isinstance(noted, str) else ""
        rate = rate_of(ccy)

        if any(h in name.lower() for h in INTERNAL_HINTS):
            internal_rows.append((name, inv_no, ccy, bal, round(bal * rate) if rate else None))
            continue

        ccy_count[cid][ccy] += 1
        if issue and due:
            terms[cid].append((datetime.fromisoformat(due) - datetime.fromisoformat(issue)).days)
        t = tier_by_name.get(name.lower())
        if cid not in customers:
            sheet_pic = str(t["pic"]) if t and t["pic"] else ""
            pic = pic_override.get(name.lower(), sheet_pic)
            if t is None and not pic:
                unmatched_tier.add(name)
            if pic != sheet_pic and sheet_pic:
                pic_changed.add(f"{name}: {sheet_pic} -> {pic}")
            tier_txt = (t["tier"] if t else None) or (str(tier_cell) if tier_cell not in (None, "#N/A") else None)
            tier_no = int(str(tier_txt).replace("Tier", "").strip()) if tier_txt and str(tier_txt).startswith("Tier") else None
            ch = chan_by_code.get(code)
            if ch and ch["ownerCompCode"] in ENTITY_BY_OWNER_CODE:
                entity = ENTITY_BY_OWNER_CODE[ch["ownerCompCode"]]
                entity_from_ellis.add(name)
            else:
                entity = CONTROL_COMPANY_BY_CHANNEL.get(name.lower(), CONTROL_COMPANY_DEFAULT)
                entity_fallback.add(name)
            customers[cid] = {
                "customer_id": cid, "customer_name": name,
                "customer_group": tier_txt or "Untiered",
                "country": (ch["countryName"] if ch else "Unknown"), "region": (t["type"] if t else ""),
                "account_owner_id": pic.lower(), "account_owner_name": pic,
                "finance_owner": None, "contract_currency": ccy, "payment_terms_days": None, "credit_limit": None,
                "credit_status": "ACTIVE", "customer_status": "ACTIVE", "collection_status": "NORMAL", "risk_grade_manual": None,
                "preferred_contact_channel": None, "control_company": entity, "tier": tier_no,
                "data_source": ("op-workbook:Tier" if t else "op-workbook:Outstanding (not in Tier sheet)") + ("+ellis:list_channels" if ch else ""),
            }

        # ELLIS remark notes -> payments already recorded in ELLIS
        disputed = 0.0
        dispute_status, dispute_reason = "NONE", None
        last_pay_date, last_pay_amt = None, None
        if isinstance(remark, str):
            for line in remark.splitlines():
                m = NOTE_HDR.match(line.strip())
                if m:
                    d, who, amt = m.group(1), m.group(2).strip(), num(m.group(3))
                    payments.append({
                        "payment_id": f"{inv_no}:{d}:{len(payments)}", "invoice_id": inv_no, "customer_id": cid, "payment_date": d,
                        "payment_amount": amt, "payment_currency": ccy, "applied_amount": amt, "unapplied_amount": 0,
                        "payment_method": "BANK_TRANSFER", "payment_reference": f"ELLIS remark by {who}", "reconciliation_status": "APPLIED",
                        # Recorded in ELLIS. PM CNFM / bank reconciliation dates are not in the workbook => "not tracked" (= record date).
                        "confirmed_at": d, "reconciled_at": d, "data_source": "op-workbook:Remark Detail (Ellis)",
                    })
                    if last_pay_date is None or d > last_pay_date:
                        last_pay_date, last_pay_amt = d, amt
                if "disput" in line.lower():
                    dispute_status = "OPEN"
                    dispute_reason = line.strip()
                    m2 = re.search(r"[￥$]\s*([\d,]+(?:\.\d+)?)", line)
                    disputed += num(m2.group(1)) if m2 else 0.0
            activities.append({
                "activity_id": f"note:{inv_no}", "customer_id": cid, "invoice_id": inv_no, "owner": "Lina (Josh)",
                "activity_type": "NOTE", "activity_date": last_pay_date or as_of, "contact_channel": None,
                "note": remark.strip()[:2000], "promised_payment_date": None, "promised_payment_amount": None, "promised_currency": None,
                "next_action": None, "next_action_date": None, "escalation_level": 0, "completed": True,
            })
        for label, text in (("vnop", vnop), ("noted", noted)):
            if isinstance(text, str) and text.strip():
                activities.append({
                    "activity_id": f"{label}:{inv_no}", "customer_id": cid, "invoice_id": inv_no, "owner": "VN OP",
                    "activity_type": "NOTE", "activity_date": as_of, "contact_channel": None, "note": text.strip(),
                    "promised_payment_date": None, "promised_payment_amount": None, "promised_currency": None,
                    "next_action": text.strip() if label == "vnop" else None, "next_action_date": None, "escalation_level": 0, "completed": label != "vnop",
                })

        # "Payment received ..." => money at the bank, ELLIS not updated yet (stage RECEIVED of the reflection chain)
        ledger_jpy += bal * (rate or 0)
        ledger_n += 1
        received = noted.lower().startswith("payment received") and bal > 0
        if received:
            reopened = "reopen" in noted.lower()
            p = prev.get(inv_no)
            if p and p["balance"] > 0 and abs(p["balance"] - bal) / max(bal, 1) < 0.05:
                pay_date, basis = as_of, "received since the previous report"
            elif p and p["last_payment"]:
                pay_date, basis = p["last_payment"], "date of the earlier ELLIS record (invoice reopened)"
            elif due and due <= as_of:
                pay_date, basis = due, "date not stated: assumed on the due date"
            else:
                pay_date, basis = as_of, "date not stated"
            payments.append({
                "payment_id": f"{inv_no}:received:{len(payments)}", "invoice_id": inv_no, "customer_id": cid, "payment_date": pay_date,
                "payment_amount": round(bal, 2), "payment_currency": ccy, "applied_amount": round(bal, 2), "unapplied_amount": 0,
                "payment_method": "BANK_TRANSFER", "payment_reference": f"OP note: {noted} ({basis})", "reconciliation_status": "APPLIED",
                "recorded_at": None, "confirmed_at": None, "reconciled_at": None,
                "data_source": "op-workbook:Noted (reopened for rate change)" if reopened else "op-workbook:Noted (awaiting AC update)",
            })
            pending_jpy += bal * (rate or 0)
            pending_n += 1
            if reopened:
                reopen_jpy += bal * (rate or 0)
                reopen_n += 1
            if last_pay_date is None or pay_date > last_pay_date:
                last_pay_date, last_pay_amt = pay_date, round(bal, 2)
            paid_t, bal_t, status = round(total, 2), 0.0, "PAID"
        else:
            paid_t, bal_t = round(paid, 2), round(bal, 2)
            status = "DISPUTED" if dispute_status == "OPEN" and abs(disputed - bal) < 0.005 else ("PARTIALLY_PAID" if paid > 0 else "OPEN")
        invoices.append({
            "invoice_id": inv_no, "invoice_number": inv_no, "booking_id": None, "customer_id": cid,
            "invoice_date": issue, "service_date": None, "due_date": due,
            "original_amount": round(total, 2), "paid_amount": paid_t, "credit_note_amount": 0,
            "disputed_amount": round(min(disputed, bal_t), 2), "outstanding_amount": bal_t,
            "invoice_currency": ccy, "invoice_status": status, "dispute_status": dispute_status if bal_t > 0 else "NONE", "dispute_reason": dispute_reason if bal_t > 0 else None,
            "cancellation_status": "NONE", "last_payment_date": last_pay_date, "last_payment_amount": last_pay_amt,
            "data_source": "op-workbook:" + cur_tab + (" (ELLIS balance %s %.2f; payment received per OP note)" % (ccy, bal) if received else ""),
        })

    for cid, c in customers.items():
        c["contract_currency"] = max(ccy_count[cid].items(), key=lambda kv: kv[1])[0]
        if terms[cid]:
            c["payment_terms_days"] = int(statistics.median(terms[cid]))

    # ---- notes ----
    notes.append(f"ELLIS ledger outstanding: JPY {ledger_jpy:,.0f} ({ledger_n} invoices). Reported by OP as payment received but not yet updated in ELLIS (AC team): JPY {pending_jpy:,.0f} ({pending_n} invoices; {reopen_n} of them reopened for a rate change, JPY {reopen_jpy:,.0f}). Open customer receivables after those receipts: JPY {ledger_jpy - pending_jpy:,.0f} ({ledger_n - pending_n} invoices).")
    if pending_n:
        notes.append("Receipts taken from the OP 'Noted' column are assumed to cover the full ELLIS balance of the invoice; amounts become exact once AC updates ELLIS.")
    if internal_rows:
        by_name = {}
        for name, _, ccy, bal, jpy in internal_rows:
            e = by_name.setdefault(name, {"n": 0, "jpy": 0})
            e["n"] += 1
            e["jpy"] += jpy or 0
        notes.append("Internal accounts excluded (not customer receivables; 출장/하드블럭 미판매분): " + ", ".join(f"{n} {e['n']} inv. JPY {e['jpy']:,.0f}" for n, e in by_name.items()) + f" — total JPY {sum(e['jpy'] for e in by_name.values()):,.0f}.")
    if unmatched_tier:
        notes.append("Sellers without an account owner (not in the Tier sheet, no override): " + ", ".join(sorted(unmatched_tier)) + ".")
    if pic_changed:
        notes.append(f"Account owners follow the assignments confirmed on {overrides.get('confirmed', '?')} (automation/config/pic-overrides.json); the workbook Tier sheet still shows the earlier PIC for: " + "; ".join(sorted(pic_changed)) + ".")
    if entity_from_ellis:
        notes.append("Managing entity and country come from ELLIS list_channels (ownerCompName / countryName)" + (f"; fallback mapping used for: {', '.join(sorted(entity_fallback))}." if entity_fallback else "."))
    else:
        notes.append("Managing entity assigned from the tracker mapping (Ctrip Vietnam -> OMH Vietnam; Ctrip Korea, Agoda, Kakao -> OMH Korea; others -> OMH Singapore HQ); no ELLIS channel file.")
    notes.append("PM CNFM and bank reconciliation dates are not in the workbook: payments already recorded in ELLIS are shown as verified/reconciled on their record date (not tracked).")
    notes.append("Credit limits and booking context are not in the workbook.")

    ds = {
        "as_of": f"{as_of}T00:00:00+07:00", "source": "file", "reporting_currency": "JPY", "fx": fx,
        "customers": list(customers.values()), "invoices": invoices, "payments": payments, "activities": activities, "bookings": [],
        "completeness": {"customers": "partial", "invoices": "full", "payments": "partial", "activities": "partial", "fx": "partial", "notes": notes},
    }
    with open(out, "w", encoding="utf-8") as f:
        json.dump(ds, f, ensure_ascii=False, indent=1)
    print(json.dumps({"sheet": cur_tab, "previous_sheet": prev_tab, "customers": len(customers), "invoices": len(invoices), "payments": len(payments), "pending_receipts": pending_n, "activities": len(activities), "fx": [r["currency"] for r in rates], "notes": notes}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else "automation/input/ellis-export.json")
