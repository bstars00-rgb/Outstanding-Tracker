"""
Build the weekly "Outstanding_Review_<date>.xlsx" from the pipeline outputs (tracker-model.json + insight.json):
  Summary        – KPIs, SOP urgency totals (formulas over the Invoices sheet), CEO decisions, owner actions
  Invoices (SOP) – every open invoice with SOP level / approval route / Tier deadline / recommended action
  Payments       – payments recorded this week (from the ELLIS remark notes)
Usage: python automation/tools/build_review_xlsx.py automation/out <report_date> <out.xlsx>
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
LEVEL_FILL = {"L1": "F8D7DA", "L2": "FFE5B4", "L3": "FFF9C4", "L4": "E8F0FE"}
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main(out_dir, report_date, out_path):
    m = json.load(open(f"{out_dir}/tracker-model.json", encoding="utf-8"))
    ins = json.load(open(f"{out_dir}/insight.json", encoding="utf-8"))["output"]
    ccy = m["reporting_currency"]
    fmt_money = '#,##0;(#,##0);"-"'
    wb = Workbook()

    # ---------------- Invoices (SOP) ----------------
    wi = wb.active
    wi.title = "Invoices (SOP)"
    hdr = ["Invoice No", "Customer", "PIC (OMH)", "Tier", "Currency", "Outstanding (orig)", f"Outstanding ({ccy})", "Issue date", "Due date", "Aging days", "Aging bucket", "SOP level", "Approval route", "Tier deadline", "Past deadline", "SOP rule", "Dispute", "Last payment", "Recommended action (tracker)", "OP decision / note"]
    wi.append(hdr)
    style_header(wi, 1, len(hdr))
    actions_by_inv = {}
    for a in m["actions"]:
        if a.get("invoice_id"):
            actions_by_inv.setdefault(a["invoice_id"], []).append(a)
    rows = [i for i in m["invoices"] if i["outstanding_amount"] > 0]
    order = {"L1": 0, "L2": 1, "L3": 2, "L4": 3, None: 4}
    rows.sort(key=lambda i: (order[i["sop"]["level"]], -i["outstanding_reporting"]))
    for i in rows:
        s = i["sop"]
        acts = sorted(actions_by_inv.get(i["invoice_id"], []), key=lambda a: {"critical": 0, "high": 1, "medium": 2, "low": 3}[a["severity"]])
        wi.append([
            i["invoice_number"], i["customer_name"], i["account_owner_name"] or "", f"Tier {s['tier']}" if s["tier"] else "Untiered", i["invoice_currency"],
            i["outstanding_amount"], i["outstanding_reporting"], i["invoice_date"], i["due_date"], i["aging_days"], i["aging_bucket"],
            s["level"] or "", {"CEO": "CEO approval (> ¥500K)", "LOCAL_DIRECTOR": "Local Director (≤ ¥500K)"}.get(s["route"], ""), s["tier_deadline"] or "", "YES" if s["past_tier_deadline"] else "",
            s["rule"], i["dispute_status"] if i["dispute_status"] != "NONE" else "", i["last_payment_date"] or "",
            "; ".join(a["recommended_action"] for a in acts[:2]), "",
        ])
    n = len(rows)
    for r in range(2, n + 2):
        lvl = wi.cell(row=r, column=12).value
        if lvl in LEVEL_FILL:
            wi.cell(row=r, column=12).fill = PatternFill("solid", fgColor=LEVEL_FILL[lvl])
        for c in (6, 7):
            wi.cell(row=r, column=c).number_format = fmt_money
        for c in range(1, len(hdr) + 1):
            wi.cell(row=r, column=c).font = Font(name=FONT)
            wi.cell(row=r, column=c).border = BORDER
        wi.cell(row=r, column=19).alignment = Alignment(wrap_text=True, vertical="top")
    wi.cell(row=n + 2, column=1, value="Total").font = Font(name=FONT, bold=True)
    wi.cell(row=n + 2, column=7, value=f"=SUM(G2:G{n + 1})").number_format = fmt_money
    wi.cell(row=n + 2, column=7).font = Font(name=FONT, bold=True)
    wi.freeze_panes = "C2"
    wi.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{n + 1}"
    autosize(wi, [11, 24, 11, 9, 9, 16, 16, 11, 11, 10, 11, 9, 24, 12, 12, 40, 10, 12, 60, 30])
    wi.cell(row=1, column=20).fill = PatternFill("solid", fgColor="FFFF00")
    wi.cell(row=1, column=20).font = Font(name=FONT, bold=True)
    wi.cell(row=1, column=20).comment = Comment("Yellow = fill in by OP (decision, promise date, note). Everything else is computed by the tracker.", "Outstanding Tracker")

    # ---------------- Summary ----------------
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = f"Weekly Outstanding Review — {report_date}"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = f"Source: OP workbook Outstanding_Report_{report_date}.xlsx → Outstanding Tracker (reporting currency {ccy}). FX: workbook approximate rates."
    ws["A2"].font = Font(name=FONT, italic=True, color="555555")
    r = 4
    ws.cell(row=r, column=1, value="KPI").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=2, value=f"Value ({ccy})").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=3, value="Interpretation").font = Font(name=FONT, bold=True)
    style_header(ws, r, 3)
    for k in m["kpis"]:
        r += 1
        ws.cell(row=r, column=1, value=k["label"]).font = Font(name=FONT)
        v = ws.cell(row=r, column=2, value=k["value"])
        v.number_format = "0.0%" if k["unit"] == "ratio" else (fmt_money if k["unit"] == "currency" else "0")
        v.font = Font(name=FONT)
        ws.cell(row=r, column=3, value=k["interpretation"]).font = Font(name=FONT)
    # Total outstanding cross-check by formula over the invoice sheet
    r += 1
    ws.cell(row=r, column=1, value="Check: sum of open invoices (formula)").font = Font(name=FONT, italic=True)
    c = ws.cell(row=r, column=2, value=f"=SUM('Invoices (SOP)'!G2:G{n + 1})")
    c.number_format, c.font = fmt_money, Font(name=FONT, italic=True)

    r += 2
    ws.cell(row=r, column=1, value="SOP urgency (overdue balance)").font = Font(name=FONT, bold=True)
    r += 1
    for col, h in enumerate(["Level", "Rule", "Invoices", f"Amount ({ccy})", "Action per SOP"], start=1):
        ws.cell(row=r, column=col, value=h)
    style_header(ws, r, 5)
    sop_text = {
        "L1": ("Tier 1 and ≥ ¥1M", "OP → Director → Finance → CEO within 24 hours"),
        "L2": ("Tier 1 < ¥1M · Tier 2 > ¥500K (assumed)", "Within 3 business days"),
        "L3": ("Tier 2 ≤ ¥500K · Tier 3 · untiered", "Weekly outstanding meeting"),
        "L4": ("< ¥100K", "Monthly review"),
    }
    for lvl in ["L1", "L2", "L3", "L4"]:
        r += 1
        ws.cell(row=r, column=1, value=lvl).fill = PatternFill("solid", fgColor=LEVEL_FILL[lvl])
        ws.cell(row=r, column=2, value=sop_text[lvl][0])
        ws.cell(row=r, column=3, value=f"=COUNTIF('Invoices (SOP)'!L2:L{n + 1},\"{lvl}\")")
        a = ws.cell(row=r, column=4, value=f"=SUMIF('Invoices (SOP)'!L2:L{n + 1},\"{lvl}\",'Invoices (SOP)'!G2:G{n + 1})")
        a.number_format = fmt_money
        ws.cell(row=r, column=5, value=sop_text[lvl][1])
        for col in range(1, 6):
            ws.cell(row=r, column=col).font = Font(name=FONT)
    r += 1
    ws.cell(row=r, column=1, value="CEO approval route (> ¥500K)").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=3, value=f"=COUNTIF('Invoices (SOP)'!M2:M{n + 1},\"CEO*\")")
    a = ws.cell(row=r, column=4, value=f"=SUMIF('Invoices (SOP)'!M2:M{n + 1},\"CEO*\",'Invoices (SOP)'!G2:G{n + 1})")
    a.number_format = fmt_money
    r += 1
    ws.cell(row=r, column=1, value="Past Tier deadline (probability checklist due)").font = Font(name=FONT, bold=True)
    ws.cell(row=r, column=3, value=f"=COUNTIF('Invoices (SOP)'!O2:O{n + 1},\"YES\")")
    a = ws.cell(row=r, column=4, value=f"=SUMIF('Invoices (SOP)'!O2:O{n + 1},\"YES\",'Invoices (SOP)'!G2:G{n + 1})")
    a.number_format = fmt_money

    r += 2
    ws.cell(row=r, column=1, value="CEO decisions required").font = Font(name=FONT, bold=True)
    r += 1
    for col, h in enumerate(["Topic", "Customer", f"Amount ({ccy})", "Recommendation", "Rationale"], start=1):
        ws.cell(row=r, column=col, value=h)
    style_header(ws, r, 5)
    for d in ins["ceo_decisions"] or [{"topic": "None this week", "customer": "", "amount": None, "recommendation": "", "rationale": ""}]:
        r += 1
        for col, v in enumerate([d["topic"], d["customer"], d["amount"], d["recommendation"], d["rationale"]], start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = Font(name=FONT)
            if col == 3:
                c.number_format = fmt_money

    r += 2
    ws.cell(row=r, column=1, value="Owner actions (this week)").font = Font(name=FONT, bold=True)
    r += 1
    for col, h in enumerate(["Owner", "Customer", f"Amount ({ccy})", "Action", "Deadline"], start=1):
        ws.cell(row=r, column=col, value=h)
    style_header(ws, r, 5)
    for a in ins["owner_actions"]:
        r += 1
        for col, v in enumerate([a["owner"], a["customer"], a["amount"], a["action"], a["deadline"]], start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = Font(name=FONT)
            if col == 3:
                c.number_format = fmt_money

    r += 2
    ws.cell(row=r, column=1, value="Data notes").font = Font(name=FONT, bold=True)
    for w in ins["data_quality_warnings"]:
        r += 1
        ws.cell(row=r, column=1, value="• " + w).font = Font(name=FONT, color="555555")
    autosize(ws, [44, 40, 14, 60, 50])

    # ---------------- Payments this week ----------------
    wp = wb.create_sheet("Payments")
    ph = ["Payment date", "Customer", "Invoice", "Currency", "Amount", f"Amount ({ccy})", "Reference", "Verified (PM CNFM)", "Reconciled (bank)"]
    wp.append(ph)
    style_header(wp, 1, len(ph))
    fx = {r_["currency"]: r_["rate_to_reporting"] for r_ in m["fx"]["rates"]}
    fx[ccy] = 1
    k = 1
    for i in m["invoices"]:
        for p in i["payments"]:
            if p["payment_date"] >= m["week"]["start"]:
                k += 1
                wp.append([p["payment_date"], i["customer_name"], i["invoice_number"], p["payment_currency"], p["payment_amount"], None, p.get("payment_reference") or "", p["confirmed_at"] or "", p["reconciled_at"] or ""])
                wp.cell(row=k, column=6, value=f"=E{k}*{fx.get(p['payment_currency'], 0)}").number_format = fmt_money
                wp.cell(row=k, column=5).number_format = fmt_money
                for c in range(1, len(ph) + 1):
                    wp.cell(row=k, column=c).font = Font(name=FONT)
    wp.cell(row=k + 1, column=1, value="Total").font = Font(name=FONT, bold=True)
    wp.cell(row=k + 1, column=6, value=f"=SUM(F2:F{k})").number_format = fmt_money
    wp.cell(row=1, column=6).comment = Comment("Converted with the workbook 'information' sheet rates (USD 150, KRW 0.11, VND 0.006 → JPY).", "Outstanding Tracker")
    autosize(wp, [13, 24, 10, 9, 16, 16, 36, 18, 18])

    # ---------------- ELLIS reflection pending (received at the bank, not yet in ELLIS) ----------------
    wr_ = wb.create_sheet("ELLIS pending")
    rh = ["Invoice No", "Customer", "PIC (OMH)", "Entity", "Currency", "Amount received", f"Amount ({ccy})", "Payment date (basis)", "Days waiting", "Record SLA (days)", "Over SLA", "Type", "Next owner", "AC update done (date)"]
    wr_.append(rh)
    style_header(wr_, 1, len(rh))
    inv_by_id = {i["invoice_id"]: i for i in m["invoices"]}
    cust_by_id = {c["customer_id"]: c for c in m["customers"]}
    q = 1
    for it in m["reflection_queue"]:
        if it["stage"] != "RECEIVED":
            continue
        q += 1
        inv = inv_by_id.get(it.get("invoice_id") or "", {})
        pay = next((p for p in inv.get("payments", []) if p["payment_id"] == it["payment_id"]), {})
        kind = "Reopened for rate change" if "reopened" in (pay.get("data_source") or "") else "Awaiting AC update"
        wr_.append([inv.get("invoice_number", ""), it["customer_name"], inv.get("account_owner_name", ""), it.get("control_company") or "", it["currency"], it["amount"], None, it["payment_date"], it["days_in_stage"], it["sla_days"], "YES" if it["overdue_sla"] else "", kind, it["next_owner"], ""])
        wr_.cell(row=q, column=7, value=f"=F{q}*{fx.get(it['currency'], 0)}").number_format = fmt_money
        wr_.cell(row=q, column=6).number_format = fmt_money
        for c in range(1, len(rh) + 1):
            wr_.cell(row=q, column=c).font = Font(name=FONT)
            wr_.cell(row=q, column=c).border = BORDER
        if it["overdue_sla"]:
            wr_.cell(row=q, column=11).fill = PatternFill("solid", fgColor=LEVEL_FILL["L1"])
    wr_.cell(row=q + 1, column=1, value="Total").font = Font(name=FONT, bold=True)
    tot = wr_.cell(row=q + 1, column=7, value=f"=SUM(G2:G{q})" if q > 1 else 0)
    tot.number_format, tot.font = fmt_money, Font(name=FONT, bold=True)
    wr_.cell(row=1, column=14).fill = PatternFill("solid", fgColor="FFFF00")
    wr_.cell(row=1, column=14).font = Font(name=FONT, bold=True)
    wr_.cell(row=1, column=8).comment = Comment("From the OP 'Noted' column. The note carries no date: received-this-week items use the report date, reopened invoices the date of the earlier ELLIS record, others the due date (assumed).", "Outstanding Tracker")
    wr_.freeze_panes = "C2"
    autosize(wr_, [11, 24, 11, 15, 9, 17, 17, 18, 12, 12, 9, 26, 16, 22])

    # ledger reconciliation on the Summary sheet
    r += 2
    ws.cell(row=r, column=1, value="ELLIS ledger vs. real receivables").font = Font(name=FONT, bold=True)
    r += 1
    for col_, h in enumerate(["Line", f"Amount ({ccy})", "Invoices", "Meaning"], start=1):
        ws.cell(row=r, column=col_, value=h)
    style_header(ws, r, 4)
    open_row, pend_row = r + 1, r + 2
    lines = [
        ("Open customer receivables (tracker)", f"=SUM('Invoices (SOP)'!G2:G{n + 1})", n, "Still unpaid by the customer"),
        ("Received at the bank, not yet in ELLIS", (f"=SUM('ELLIS pending'!G2:G{q})" if q > 1 else 0), q - 1, "AC team must record / re-confirm in ELLIS"),
        ("ELLIS ledger outstanding (open + pending)", f"=B{open_row}+B{pend_row}", f"=C{open_row}+C{pend_row}", "What ELLIS shows today"),
    ]
    for label, amount, count, meaning in lines:
        r += 1
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, bold=label.startswith("ELLIS ledger"))
        c = ws.cell(row=r, column=2, value=amount)
        c.number_format, c.font = fmt_money, Font(name=FONT, bold=label.startswith("ELLIS ledger"))
        ws.cell(row=r, column=3, value=count).font = Font(name=FONT)
        ws.cell(row=r, column=4, value=meaning).font = Font(name=FONT, color="555555")

    wb.save(out_path)
    print("saved", out_path, "invoices", n, "payments", k - 1, "ellis_pending", q - 1)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
